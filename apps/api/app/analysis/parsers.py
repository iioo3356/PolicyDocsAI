import csv
import io
import re
from tree_sitter_language_pack import get_parser


from app.domain.parsed_chunk import ParsedChunk
from app.domain.business_policy_signal import business_policy_score, is_business_policy_code


POLICY_MARKERS = re.compile(
    r"\b(if|else if|switch|when|throw|require|check|validate|permission|disabled|"
    r"can[A-Z_]|isVisible|status|limit|max|min|amount|point|cancel|apply|approve)\b|"
    r"(권한|상태|취소|지원|신청|승인|포인트|금액|기간|가능|불가)",
    re.IGNORECASE,
)
LANGUAGE_BY_SUFFIX = {".ts": "typescript", ".tsx": "tsx", ".js": "javascript", ".jsx": "tsx", ".kt": "kotlin"}
DECISION_NODES = {"if_statement", "switch_statement", "ternary_expression", "if_expression", "when_expression"}
VALIDATION_CALL = re.compile(r"^(require|check|validate|assert|invariant)$", re.IGNORECASE)


def parse_markdown(text: str) -> list[ParsedChunk]:
    lines = text.splitlines()
    chunks: list[ParsedChunk] = []
    headings: list[str] = []
    start = 1
    buffer: list[str] = []

    def flush(end: int) -> None:
        nonlocal buffer, start
        content = "\n".join(buffer).strip()
        if content:
            chunks.append(ParsedChunk("heading", content, start, end, " > ".join(headings), score=0.75))
        buffer = []

    for number, line in enumerate(lines, 1):
        match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if match:
            flush(number - 1)
            level = len(match.group(1))
            headings[:] = headings[: level - 1]
            headings.append(match.group(2).strip())
            start = number
        buffer.append(line)
        if len(buffer) >= 80:
            flush(number)
            start = number + 1
    flush(len(lines))
    return chunks


def parse_csv(text: str) -> list[ParsedChunk]:
    rows = list(csv.DictReader(io.StringIO(text)))
    result = []
    for index, row in enumerate(rows, 2):
        cleaned = {key: value for key, value in row.items() if key and value and value.strip()}
        if cleaned:
            content = "\n".join(f"{key}: {value}" for key, value in cleaned.items())
            result.append(ParsedChunk("csv_row", content, index, index, f"row {index}", score=0.7, metadata=cleaned))
    return result


def parse_xlsx(raw: bytes) -> list[ParsedChunk]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    result: list[ParsedChunk] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            first = next(rows, None)
            if first is None:
                continue
            headers = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(first)]
            for row_number, row in enumerate(rows, 2):
                values = []
                metadata = {}
                for index, value in enumerate(row):
                    if value is None or str(value).strip() == "":
                        continue
                    header = headers[index] if index < len(headers) else f"column_{index + 1}"
                    rendered = str(value).strip()
                    values.append(f"{header}: {rendered}")
                    metadata[header] = rendered
                if values:
                    result.append(ParsedChunk(
                        "xlsx_row", "\n".join(values), row_number, row_number,
                        f"{sheet.title} > row {row_number}", score=0.7, metadata=metadata,
                    ))
    finally:
        workbook.close()
    return result


def _decision_blocks(text: str, suffix: str) -> list[tuple[int, int, str]]:
    source = text.encode("utf-8")
    parser = get_parser(LANGUAGE_BY_SUFFIX.get(suffix, "typescript"))
    root = parser.parse(source).root_node
    blocks = []
    pending = [root]
    while pending:
        node = pending.pop()
        if node.type in DECISION_NODES:
            blocks.append((node.start_point[0] + 1, node.end_point[0] + 1,
                           source[node.start_byte:node.end_byte].decode("utf-8", errors="replace")))
            continue
        if node.type == "call_expression":
            function = node.child_by_field_name("function")
            name = source[function.start_byte:function.end_byte].decode() if function else ""
            if VALIDATION_CALL.fullmatch(name.rsplit(".", 1)[-1]):
                parent = node.parent if node.parent and node.parent.type == "expression_statement" else node
                blocks.append((parent.start_point[0] + 1, parent.end_point[0] + 1,
                               source[parent.start_byte:parent.end_byte].decode("utf-8", errors="replace")))
                continue
        pending.extend(reversed(node.children))
    return blocks


def parse_code(text: str, suffix: str = ".ts") -> list[ParsedChunk]:
    chunks = []
    for start_line, end_line, content in _decision_blocks(text, suffix):
        if not is_business_policy_code(content):
            continue
        marker_count = len(POLICY_MARKERS.findall(content))
        signal_score = business_policy_score(content)
        score = min(0.95, 0.48 + marker_count * 0.05 + signal_score * 0.04)
        chunks.append(ParsedChunk("code", content.strip(), start_line, end_line, score=score,
                                  metadata={"business_signal_score": signal_score}))
    return chunks
